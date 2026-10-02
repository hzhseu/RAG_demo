import json
import zipfile
from pathlib import Path
import numpy as np
import pytest

from nordrag.parsing import parse_pptx, chunk_pages, scan
from nordrag.package import seal, open_package, export_package, PackageError
from nordrag.index import create_index, retrieve


def test_extraction_preserves_page_table_units_and_notes(deck):
    pages = parse_pptx(deck)
    assert len(pages) == 2
    assert any(b["kind"] == "notes" and "support" in b["text"] for b in pages[0]["blocks"])
    table = next(b for b in pages[1]["blocks"] if b["kind"] == "table")
    assert table["rows"][1] == ["Alpha", "120.50", "12.5"]
    chunks = chunk_pages("doc", pages, max_chars=90)
    rows = [c for c in chunks if "120.50" in c["text"]]
    assert rows and rows[0]["page"] == 2 and "EUR" in rows[0]["text"]


def test_scan_ignores_temporary_and_deduplicates(deck):
    (deck.parent / "copy.pptx").write_bytes(deck.read_bytes())
    (deck.parent / "~$temp.pptx").write_bytes(b"invalid")
    result = scan(deck.parent)
    assert len(result["documents"]) == 1
    assert len(result["duplicates"]) == 1
    assert not result["failed"]


def make_root(tmp_path):
    root = tmp_path / "content"
    root.mkdir()
    (root / "doc.txt").write_text("Alpha 120.50 EUR", encoding="utf-8")
    (root / "documents.json").write_text(json.dumps([{"id": "a", "tags": []}]), encoding="utf-8")
    return root


def test_package_portable_export_updates_identity(tmp_path):
    root = make_root(tmp_path)
    p = tmp_path / "first.ragkb"
    seal(root, p, {"name": "部门", "embedding": "test-embedding"})
    work, first = open_package(p, tmp_path / "work", "test-embedding")
    assert (work / "doc.txt").read_text() == "Alpha 120.50 EUR"
    docs = json.loads((work / "documents.json").read_text(encoding="utf-8"))
    docs[0]["tags"] = ["财务"]
    out = tmp_path / "second.ragkb"
    export_package(work, out, docs)
    second_work, second = open_package(out, tmp_path / "second", "test-embedding")
    assert first["version"] != second["version"]
    assert json.loads((second_work / "documents.json").read_text(encoding="utf-8"))[0]["tags"] == ["财务"]
    with pytest.raises(FileExistsError):
        seal(root, p, {"embedding": "test-embedding"})


def test_package_rejects_tampering_mismatch_and_traversal(tmp_path):
    root = make_root(tmp_path)
    p = tmp_path / "good.ragkb"
    seal(root, p, {"embedding": "expected"})
    with pytest.raises(PackageError, match="模型"):
        open_package(p, tmp_path / "mismatch", "wrong")
    with zipfile.ZipFile(p, "a") as z:
        z.writestr("../escape.txt", "evil")
    with pytest.raises(PackageError):
        open_package(p, tmp_path / "bad", "expected")
    assert not (tmp_path / "escape.txt").exists()


def test_hybrid_retrieves_cjk_and_preserves_source(tmp_path):
    chunks = [{"id": "a", "doc_id": "d", "page": 2, "kind": "table", "text": "产品收入 Alpha 120.50 EUR"}, {"id": "b", "doc_id": "e", "page": 1, "kind": "text", "text": "support warranty"}]
    create_index(tmp_path, chunks, np.array([[1., 0.], [0., 1.]]))
    hits = retrieve(tmp_path, "收入", np.array([1., 0.]))
    assert hits[0]["id"] == "a" and hits[0]["page"] == 2
    assert hits[0]["text"].endswith("120.50 EUR")


def test_merged_table_cell_is_not_duplicated(deck):
    from pptx import Presentation
    prs=Presentation(deck)
    table=prs.slides[1].shapes[0].table
    table.cell(1,0).merge(table.cell(2,0))
    prs.save(deck)
    block=next(b for b in parse_pptx(deck)[1]['blocks'] if b['kind']=='table')
    assert 'Alpha' in block['rows'][1][0]
    assert block['rows'][2][0]==''
    assert block['rows'][2][1]=='99.00'


def test_valid_manifest_cannot_hide_changed_payload(tmp_path):
    root=make_root(tmp_path)
    original=tmp_path/'original.ragkb'; damaged=tmp_path/'damaged.ragkb'
    seal(root,original,{'embedding':'e'})
    with zipfile.ZipFile(original) as src, zipfile.ZipFile(damaged,'w') as dst:
        for name in src.namelist():
            dst.writestr(name,b'changed' if name=='doc.txt' else src.read(name))
    with pytest.raises(PackageError,match='校验失败'):
        open_package(damaged,tmp_path/'work','e')


def test_page_context_keeps_headings_with_facts_and_table_year():
    pages=[{'page':1,'blocks':[{'kind':'text','text':'Approval workflow'},{'kind':'text','text':'Above EUR 5000 requires manager approval.'}]}, {'page':2,'blocks':[{'kind':'text','text':'Revenue 2025'},{'kind':'table','text':'','rows':[['Product','EUR'],['Alpha','120.50']]}]}]
    chunks=chunk_pages('d',pages)
    approval=next(c for c in chunks if 'Approval workflow' in c['text'])
    assert '5000' in approval['text']
    revenue=next(c for c in chunks if '120.50' in c['text'])
    assert '2025' in revenue['text']
