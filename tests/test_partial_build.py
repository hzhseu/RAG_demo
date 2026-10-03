import threading

import numpy as np
import pytest
from pptx import Presentation
from pptx.util import Inches

from nordrag import builder
from nordrag.package import open_package
from nordrag.index import all_chunks, retrieve
from nordrag.util import read_json
from test_build import Engines


@pytest.fixture
def pair(deck):
    bad = deck.parent / 'bad.pptx'
    prs = Presentation()
    for i in range(10):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = f'BAD document page {i}'
    prs.save(bad)
    return deck, bad


def run(tmp_path, engines, name='result', cancel=None):
    return builder.build(tmp_path, 'Partial', tmp_path / f'{name}.ragkb', tmp_path / 'cache', engines, cancel or threading.Event())


def assert_partial(tmp_path, good, result):
    assert result['status'] == 'complete_with_warnings'
    assert result['included'] == [good.name]
    assert result['total_pages'] == 2
    root, manifest = open_package(tmp_path / 'result.ragkb', tmp_path / 'work', Engines.embedding_id)
    docs = read_json(root / 'documents.json')
    assert len(docs) == manifest['documents'] == 1
    assert manifest['pages'] == 2
    assert {p.name for p in (root / 'documents').iterdir()} == {docs[0]['id']}
    chunks = read_json(root / 'chunks.json')
    assert chunks == all_chunks(root)
    assert len(chunks) == len(np.load(root / 'vectors.npy')) == result['chunks']
    assert all(c['doc_id'] == docs[0]['id'] for c in chunks)
    assert all(c['doc_id'] == docs[0]['id'] for c in docs[0]['citations'])
    assert retrieve(root, 'Alpha', [1, 0, 0])
    assert read_json(root / 'build-report.json')['failed'] == result['failed']
    assert not list((tmp_path / 'cache').glob('build-*'))


@pytest.mark.parametrize('stage', ['convert', 'organize', 'embedding'])
def test_partial_package_has_no_failed_document_residue(pair, tmp_path, stage):
    good, bad = pair
    e = Engines()
    original = getattr(e, 'embed' if stage == 'embedding' else stage)
    bad_batches = 0
    def fail(*args):
        nonlocal bad_batches
        is_bad = (args[0].name == bad.name if stage == 'convert' else
                  any('BAD' in (c if stage == 'embedding' else c['text']) for c in args[0]))
        if is_bad:
            bad_batches += 1
            if stage != 'embedding' or bad_batches == 2:
                raise TimeoutError('document timeout')
        return original(*args)
    setattr(e, 'embed' if stage == 'embedding' else stage, fail)
    result = run(tmp_path, e)
    assert result['failed'][0]['path'] == bad.name
    assert result['failed'][0]['stage'] == stage
    assert_partial(tmp_path, good, result)


def test_fixed_document_rejoins_and_successful_cache_is_reused(pair, tmp_path):
    good, bad = pair
    e = Engines()
    original = e.organize
    def organize(chunks, cancel):
        if any('BAD' in c['text'] for c in chunks):
            raise ValueError('invalid summary')
        return original(chunks, cancel)
    e.organize = organize
    run(tmp_path, e)
    e.convert = lambda *a: pytest.fail('parse cache was not reused')
    def repaired(chunks, cancel):
        assert any('BAD' in c['text'] for c in chunks), 'successful summary was not reused'
        return original(chunks, cancel)
    e.organize = repaired
    result = run(tmp_path, e, 'fixed')
    assert result['status'] == 'complete'
    assert len(result['included']) == 2 and result['total_pages'] == 12


@pytest.mark.parametrize('target', ['write_json', 'create_index', 'seal'])
def test_shared_storage_and_publication_errors_abort(pair, tmp_path, monkeypatch, target):
    original = getattr(builder, target)
    def fail(*args, **kwargs):
        if target != 'write_json' or args[0].name == 'organization.json':
            raise PermissionError('shared write denied')
        return original(*args, **kwargs)
    monkeypatch.setattr(builder, target, fail)
    with pytest.raises(builder.BuildError, match='shared write denied'):
        run(tmp_path, Engines())
    assert not (tmp_path / 'result.ragkb').exists()


def test_cancel_during_embedding_does_not_publish(pair, tmp_path):
    cancel = threading.Event()
    e = Engines()
    def embed(texts):
        cancel.set()
        return np.array([[1., 0., 0.] for _ in texts])
    e.embed = embed
    with pytest.raises(builder.BuildError, match='取消'):
        run(tmp_path, e, cancel=cancel)
    assert not (tmp_path / 'result.ragkb').exists()


@pytest.mark.parametrize('stage', ['organize', 'embed'])
def test_lazy_shared_model_start_failure_is_not_skipped(pair, tmp_path, stage):
    from nordrag.diagnostics import ComponentUnavailableError
    e = Engines()
    def fail(*args):
        raise ComponentUnavailableError('model cannot start')
    setattr(e, stage, fail)
    with pytest.raises(builder.BuildError, match='model cannot start') as error:
        run(tmp_path, e)
    assert error.value.report['failed'] == []
    assert not (tmp_path / 'result.ragkb').exists()


@pytest.mark.parametrize('kind', ['empty', 'nan', 'zero', 'dimension'])
def test_invalid_document_vectors_are_excluded(pair, tmp_path, kind):
    good, bad = pair
    e = Engines()
    original = e.embed
    def embed(texts):
        if any('BAD' in t for t in texts):
            if kind == 'empty':
                return []
            if kind == 'nan':
                return [[float('nan'), 0, 0] for t in texts]
            if kind == 'zero':
                return [[0, 0, 0] for t in texts]
            # Mismatched dimensions within one document, across batch boundary.
            return [[1, 0] if 'page 9' in t else [1, 0, 0] for t in texts]
        return original(texts)
    e.embed = embed
    result = run(tmp_path, e)
    assert_partial(tmp_path, good, result)


def test_parse_failure_is_skipped(pair, tmp_path, monkeypatch):
    good, bad = pair
    original = builder.parse_pptx
    def parse(path, *args):
        if path == bad:
            raise ValueError('unreadable slide image')
        return original(path, *args)
    monkeypatch.setattr(builder, 'parse_pptx', parse)
    result = run(tmp_path, Engines())
    assert result['failed'][0]['stage'] == 'parse'
    assert_partial(tmp_path, good, result)


@pytest.mark.parametrize('mode', ['empty', 'all_failed'])
def test_no_usable_documents_never_publish(tmp_path, mode):
    if mode == 'all_failed':
        (tmp_path / 'broken.pptx').write_bytes(b'broken')
    with pytest.raises(builder.BuildError, match='未生成知识库'):
        run(tmp_path, Engines())
    assert not (tmp_path / 'result.ragkb').exists()


def test_corrupt_embedded_image_is_a_document_failure(deck, tmp_path):
    from io import BytesIO
    from zipfile import ZipFile
    from PIL import Image
    image = BytesIO()
    Image.new('RGB', (8, 8), 'white').save(image, format='PNG')
    image.seek(0)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(image, Inches(1), Inches(1))
    original = BytesIO()
    prs.save(original)
    with ZipFile(original) as source, ZipFile(tmp_path / 'bad-image.pptx', 'w') as target:
        for name in source.namelist():
            target.writestr(name, b'corrupt image' if name.startswith('ppt/media/') else source.read(name))
    result = run(tmp_path, Engines())
    assert result['failed'][0]['path'] == 'bad-image.pptx'
    assert result['failed'][0]['stage'] == 'parse'
    assert_partial(tmp_path, deck, result)


def test_wrong_batch_vector_counts_cannot_shift_chunk_alignment(pair, tmp_path):
    good, bad = pair
    e = Engines()
    original = e.embed
    calls = 0
    def embed(texts):
        nonlocal calls
        if any('BAD' in t for t in texts):
            calls += 1
            count = len(texts) + (-1 if calls == 1 else 1)
            return [[1, 0, 0] for _ in range(count)]
        return original(texts)
    e.embed = embed
    result = run(tmp_path, e)
    assert_partial(tmp_path, good, result)
