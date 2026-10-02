from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches


@pytest.fixture
def deck(tmp_path: Path):
    path = tmp_path / "中英 测试.pptx"
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = "Nord 产品 Alpha: warranty 保修 24 months."
    s.notes_slide.notes_text_frame.text = "Internal note: support 支持 email."
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(8), Inches(2)).table
    for r, row in enumerate([["产品 Product", "收入 Revenue (EUR)", "增长 Growth (%)"], ["Alpha", "120.50", "12.5"], ["Beta", "99.00", "8.0"]]):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    prs.save(path)
    return path
