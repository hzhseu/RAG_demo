import threading
import json
import pytest
from fastapi.testclient import TestClient
from nordrag.api import create_app, create_knowledge_app
from nordrag.builder import build
from nordrag.package import open_package
from test_build import Engines


class ChatEngines(Engines):
    def cancel(self): pass
    def count(self, text): return len(text)
    def stream(self, messages, cancel, max_tokens=768):
        yield "保修为24个月 [1]"


def setup_client(deck, tmp_path):
    engines = ChatEngines()
    out = tmp_path / "test.ragkb"
    build(deck.parent, "部门", out, tmp_path / "cache", engines, threading.Event())
    root, manifest = open_package(out, tmp_path / "work", engines.embedding_id)
    app = create_app(root, manifest, engines, tmp_path / "data", token="test-secret")
    app.state.engines = engines
    return TestClient(app), {"X-Nord-Token": "test-secret", "X-Knowledge-Sequence": "1"}


def test_api_requires_token_and_exposes_version(deck, tmp_path):
    client, headers = setup_client(deck, tmp_path)
    assert client.get("/api/status").status_code == 403
    response = client.get("/api/status", headers=headers)
    assert response.json()["knowledge"]["name"] == "部门"
    assert client.get("/api/status", headers=headers | {"Origin": "https://evil.example"}).status_code == 403


def test_stream_persists_sources_and_session_delete(deck, tmp_path):
    client, headers = setup_client(deck, tmp_path)
    session = client.post("/api/sessions", json={}, headers=headers).json()["id"]
    response = client.post("/api/chat", json={"session_id": session, "question": "保修多久？"}, headers=headers)
    assert response.status_code == 200
    assert '"type": "done"' in response.text
    history = client.get(f"/api/sessions/{session}", headers=headers).json()
    assert history["messages"][-1]["citations"][0]["page"] == 1
    assert client.delete(f"/api/sessions/{session}", headers=headers).status_code == 200
    assert client.get(f"/api/sessions/{session}", headers=headers).status_code == 404


def test_chat_counts_full_catalog_without_model_or_retrieval_despite_wrong_history(deck, tmp_path, monkeypatch):
    from pptx import Presentation
    for i in range(9):
        presentation = Presentation(deck)
        presentation.slides[0].shapes.add_textbox(0, 0, 1000000, 1000000).text = f'Project {i}'
        presentation.save(deck.parent / f'project-{i}.pptx')
    client, headers = setup_client(deck, tmp_path)
    assert len(client.get('/api/documents', headers=headers).json()) == 10
    def unexpected(*args, **kwargs):
        raise AssertionError('catalog totals must not call inference or retrieval')
    monkeypatch.setattr('nordrag.api.retrieve', unexpected)
    for method in ('embed', 'count', 'stream'):
        monkeypatch.setattr(client.app.state.engines, method, unexpected)
    sid = client.post('/api/sessions', json={}, headers=headers).json()['id']
    from nordrag.util import read_json, write_json
    path = tmp_path / 'data' / 'sessions' / f'{sid}.json'
    session = read_json(path)
    session['messages'] = [{'role': 'user', 'content': '有几个项目？'},
                           {'role': 'assistant', 'content': '共有6个项目 [项目统计]'}]
    write_json(path, session)
    for question in ('文档库中有多少文档？', '有几个项目？'):
        response = client.post('/api/chat', json={'session_id': sid, 'question': question}, headers=headers)
        events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: ')]
        done = next(e for e in events if e['type'] == 'done')
        assert '10 份 PPTX' in done['text'] and '10 个项目' in done['text']
        assert done['answer_source'] == {
            'type': 'catalog', 'scope': 'all_documents',
            'knowledge_version': session['knowledge_version'], 'document_count': 10, 'project_count': 10,
        }
        answer = client.get(f'/api/sessions/{sid}', headers=headers).json()['messages'][-1]
        assert answer['content'] == done['text']
        assert answer['answer_source'] == done['answer_source']
        assert answer['supported'] is True and answer['citations'] == []


@pytest.mark.parametrize('question', ['这个文档提到几个项目？', '涉及无线通信的项目有几个？'])
def test_scoped_count_still_uses_document_evidence(deck, tmp_path, question):
    client, headers = setup_client(deck, tmp_path)
    sid = client.post('/api/sessions', json={}, headers=headers).json()['id']
    client.post('/api/chat', json={'session_id': sid, 'question': question}, headers=headers)
    answer = client.get(f'/api/sessions/{sid}', headers=headers).json()['messages'][-1]
    assert answer['content'] == '保修为24个月 [1]'
    assert answer['citations'] and not answer.get('answer_source')


@pytest.mark.parametrize('total', [0, 1, 12])
def test_catalog_count_uses_loaded_directory_including_empty_library(tmp_path, total):
    from nordrag.util import write_json
    root = tmp_path / 'knowledge'
    write_json(root / 'documents.json', [{'id': str(i), 'name': f'{i}.pptx'} for i in range(total)])
    engine = ChatEngines()
    def unexpected(*args, **kwargs):
        raise AssertionError('catalog statistics do not require models or an index')
    engine.embed = engine.stream = engine.count = unexpected
    app = create_knowledge_app(root, {'version': f'catalog-{total}'}, engine, tmp_path / 'data', token='test')
    with TestClient(app) as client:
        headers = {'X-Nord-Token': 'test'}
        sid = client.post('/api/sessions', json={}, headers=headers).json()['id']
        response = client.post('/api/chat', json={'session_id': sid, 'question': '项目数量是多少？'}, headers=headers)
        assert '"type": "done"' in response.text
        answer = client.get(f'/api/sessions/{sid}', headers=headers).json()['messages'][-1]
        assert answer['answer_source']['document_count'] == total
        assert answer['answer_source']['project_count'] == total
        assert f'{total} 份 PPTX' in answer['content'] and f'{total} 个项目' in answer['content']


def test_content_question_without_evidence_does_not_call_model(deck, tmp_path, monkeypatch):
    client, headers = setup_client(deck, tmp_path)
    monkeypatch.setattr('nordrag.api.retrieve', lambda *args: [])
    def unexpected(*args, **kwargs):
        raise AssertionError('no evidence for content generation')
    client.app.state.engines.stream = unexpected
    sid = client.post('/api/sessions', json={}, headers=headers).json()['id']
    client.post('/api/chat', json={'session_id': sid, 'question': '项目的预算是多少？'}, headers=headers)
    answer = client.get(f'/api/sessions/{sid}', headers=headers).json()['messages'][-1]
    assert not answer['supported'] and not answer.get('answer_source')
    assert '没有可用证据' in answer['content']


def test_labels_export_roundtrip(deck, tmp_path):
    client, headers = setup_client(deck, tmp_path)
    doc = client.get("/api/documents", headers=headers).json()[0]
    assert client.patch(f"/api/documents/{doc['id']}", json={"category": "产品", "tags": ["测试"]}, headers=headers).status_code == 200
    export = client.post("/api/export", headers=headers)
    assert export.status_code == 200
    path = tmp_path / "exported.ragkb"
    path.write_bytes(export.content)
    root, _ = open_package(path, tmp_path / "export-work", "test-only-embedding")
    from nordrag.util import read_json
    assert read_json(root / "documents.json")[0]["tags"] == ["测试"]


def test_export_retains_lock_until_response_completion(deck, tmp_path):
    import asyncio
    import pytest
    from fastapi import HTTPException
    client, headers = setup_client(deck, tmp_path)
    route = lambda path: next(r.endpoint for r in client.app.state.knowledge.active.app.routes if r.path == path)
    response = route('/api/export')()
    try:
        with pytest.raises(HTTPException) as error:
            route('/api/cache/clear')()
        assert error.value.status_code == 409
    finally:
        async def send(message): pass
        async def receive(): return {"type": "http.request", "body": b""}
        asyncio.run(response({"type":"http","method":"POST","headers":[],"extensions":{}},receive,send))
    assert route('/api/cache/clear')()['cleared']


def test_invalid_export_range_cannot_leak_operation_lock(deck, tmp_path):
    client, headers = setup_client(deck, tmp_path)
    result=client.post('/api/export',headers=headers|{'Range':'bytes=999999999999-'})
    assert result.status_code==416
    assert client.post('/api/cache/clear',headers=headers).status_code==200


@pytest.mark.parametrize('mode', ['knowledge', 'advanced'])
def test_cancellation_during_embedding_does_not_load_chat(deck, tmp_path, mode):
    import concurrent.futures
    import time
    client, headers = setup_client(deck, tmp_path)
    # Engines are an external service boundary: block until cancellation is delivered.
    engine = client.app.state.engines
    entered, stopped = threading.Event(), threading.Event()
    def embed(*args, **kwargs):
        entered.set()
        if not stopped.wait(3):
            raise RuntimeError('cancel was not delivered to embedding')
        return [[1.,0.,0.]]
    engine.embed = embed
    engine.cancel = stopped.set
    engine.count = lambda text: (_ for _ in ()).throw(AssertionError('chat must not load after cancel'))
    sid = client.post('/api/sessions',headers=headers,json={'mode':mode}).json()['id']
    with concurrent.futures.ThreadPoolExecutor() as pool:
        response = pool.submit(client.post,'/api/chat',headers=headers,json={'session_id':sid,'question':'test'})
        assert entered.wait(2)
        jid = next(iter(client.app.state.knowledge.active.app.state.jobs))
        assert client.post(f'/api/jobs/{jid}/cancel',headers=headers).status_code == 200
        result = response.result(timeout=3)
    assert stopped.is_set()
    assert '"type": "cancelled"' in result.text
    assert not client.get(f'/api/sessions/{sid}',headers=headers).json()['messages']
