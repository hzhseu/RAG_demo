"""Check the bundled font can be loaded privately and embedded in a preview PDF."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from nordrag.config import load_config
from nordrag.engines import Engines
from nordrag.util import write_json

folder = ROOT / 'artifacts/font-test'
folder.mkdir(exist_ok=True)
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
run = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.paragraphs[0].add_run()
run.text = '中文知识库 Private font 120.50 EUR'
run.font.name = 'Noto Sans CJK SC'
run.font.size = Pt(28)
east = OxmlElement('a:ea')
east.set('typeface', 'Noto Sans CJK SC')
run._r.get_or_add_rPr().append(east)
source, target = folder / 'font.pptx', folder / 'font.pdf'
prs.save(source)
engine = Engines(load_config(), build=True)
try:
    engine.convert(source, target, 1)
finally:
    engine.close()
pdf = PdfReader(target)
fonts = [str(value.get_object().get('/BaseFont', '')) for value in pdf.pages[0]['/Resources']['/Font'].get_object().values()]
text = pdf.pages[0].extract_text()
assert any('NotoSansCJKsc' in font for font in fonts), fonts
assert '中文知识库' in text.replace(' ', ''), text
write_json(folder / 'report.json', {'pass': True, 'fonts': fonts, 'text': text, 'private_font_path': 'share/fonts/truetype'})
print('Bundled Noto CJK font embedded and Chinese text preserved')
